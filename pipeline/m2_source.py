"""Nguồn thật M2: parser có kiểu, audit XLSX độc lập, import bất biến."""

import argparse
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import sqlite3
import time

import openpyxl
from bson.int64 import Int64
from pymongo import MongoClient, UpdateOne

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'data/raw/online-retail-ii/online_retail_II.xlsx'
OUT = ROOT / 'artifacts/m2'
SOURCE_HASH = 'bcbe73b35f5b7babf197fb0cb983a11f5d9ff929078d4aa53d171b1f2df2e980'
HEADERS = ['Invoice', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'Price', 'Customer ID', 'Country']
SHEETS = ['Year 2009-2010', 'Year 2010-2011']
COUNTS = [525461, 541910]
SCHEMA = 'm0-contract-v1'
PARSER = 'openpyxl-3.1.5-typed-v1'
POLICY = {'version': 'p01-first-p02-test-v1', 'excluded': sorted({
    'POST', 'DOT', 'C2', 'BANK CHARGES', 'AMAZONFEE', 'D', 'S',
    'ADJUST', 'ADJUST2', 'B', 'M', 'GIFT', 'TEST001', 'TEST002'}),
    'voucher_prefix_case_insensitive': 'gift_', 'scope': 'merchandise',
    'customer_required': False, 'basket_time': 'min-local', 'single_item_baskets': 'keep'}
URI = 'mongodb://127.0.0.1:27019/?directConnection=true&serverSelectionTimeoutMS=10000'
DATABASE = 'bigdata_m2'


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def digest(value):
    return hashlib.sha256(canonical(value).encode()).hexdigest()


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def check(ok, code):
    if not ok:
        raise RuntimeError(code)


def encode_cell(value):
    if value is None:
        return {'type': 'null', 'value': None}
    if isinstance(value, datetime):
        check(value.tzinfo is None, 'UNEXPECTED_SOURCE_TIMEZONE')
        return {'type': 'local_datetime', 'value': value.isoformat()}
    if isinstance(value, bool):
        raise ValueError('UNSUPPORTED_BOOLEAN_CELL')
    if isinstance(value, (int, float)):
        number = Decimal(str(value))
        check(number.is_finite(), 'NONFINITE_SOURCE_NUMBER')
        return {'type': 'number', 'value': format(number, 'f')}
    if isinstance(value, str):
        return {'type': 'text', 'value': value}
    raise ValueError('UNSUPPORTED_SOURCE_TYPE')


def tuple_key(cells):
    values = []
    for cell in cells:
        value = cell['value']
        if cell['type'] == 'number':
            number = Decimal(value)
            value = '0' if not number else format(number.normalize(), 'f')
        values.append([cell['type'], value])
    return canonical(values)


def identifier(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = Decimal(str(value))
        return str(int(number)) if number == number.to_integral_value() else None
    return value.strip() if isinstance(value, str) else None


def excluded(code):
    return code.upper() in POLICY['excluded'] or code.lower().startswith('gift_')


def verified_rows():
    with SOURCE.open('rb') as stream:
        check(hashlib.file_digest(stream, 'sha256').hexdigest() == SOURCE_HASH, 'SOURCE_HASH')
    workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=False)
    try:
        check(workbook.sheetnames == SHEETS, 'SOURCE_SHEETS')
        for index, sheet in enumerate(workbook.worksheets):
            rows = sheet.iter_rows()
            check([cell.value for cell in next(rows)] == HEADERS, 'SOURCE_HEADER')
            count = 0
            for row_num, cells in enumerate(rows, 2):
                check(len(cells) == 8 and not any(c.data_type in ('f', 'e', 'b') for c in cells),
                      'UNSUPPORTED_FORMULA_ERROR_BOOLEAN')
                values = tuple(c.value for c in cells)
                count += 1
                yield index, row_num, values
            check(count == COUNTS[index], 'SOURCE_SHEET_COUNT')
    finally:
        workbook.close()


def document(index, row_num, values):
    cells = [encode_cell(value) for value in values]
    return {'_id': canonical([SOURCE_HASH, index, row_num]), 'dataset_id': SOURCE_HASH,
            'sheet_index': Int64(index), 'source_row': Int64(row_num),
            'source_cells': cells, 'tuple_key': tuple_key(cells)}


def position_limbs(index, row_num):
    value = hashlib.sha256(f'{index}:{row_num}'.encode()).digest()
    return [int.from_bytes(value[i:i+8], 'big') for i in range(0, 32, 8)]


def audit(output):
    """Relecture XLSX: SQLite equality independent of MongoDB/Spark/window."""
    started = time.monotonic()
    # A unique output directory prevents overwriting previous evidence/seen sets.
    output.mkdir(parents=True, exist_ok=False)
    con = sqlite3.connect(output / 'seen.sqlite')
    con.execute('CREATE TABLE seen (key TEXT PRIMARY KEY) WITHOUT ROWID')
    counts, reasons, missing = Counter(), Counter(), Counter()
    type_counts = {name: Counter() for name in HEADERS}
    scale, integer_digits, extrema = Counter(), Counter(), {}
    baskets, item_units, item_rows, item_money = {}, Counter(), Counter(), defaultdict(Decimal)
    excluded_counts, special = Counter(), {}
    reference_limbs = [0] * 4
    examples = []
    amount = Decimal(0)
    positive_amount = Decimal(0)
    per_sheet = Counter()
    normalization = {i: defaultdict(set) for i in (0, 1, 6)}
    for index, row_num, values in verified_rows():
        counts['raw'] += 1
        per_sheet[str(index)] += 1
        for name, value in zip(HEADERS, values):
            type_counts[name][type(value).__name__] += 1
            missing[name] += int(value is None or isinstance(value, str) and not value.strip())
        for col in (3, 5):
            value = values[col]
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                d = Decimal(str(value))
                s = max(0, -d.normalize().as_tuple().exponent) if d else 0
                scale[HEADERS[col]] = max(scale[HEADERS[col]], s)
                integer_digits[HEADERS[col]] = max(integer_digits[HEADERS[col]], max(1, d.adjusted()+1))
                limits = extrema.setdefault(HEADERS[col], [d, d])
                limits[0], limits[1] = min(limits[0], d), max(limits[1], d)
        # Independent representation: raw Python decoded values, tagged by category;
        # no call to encode_cell/tuple_key and no stored Mongo auxiliary key.
        independent = []
        for value in values:
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                independent.append(['n', str(Decimal(str(value)).normalize())])
            elif isinstance(value, datetime):
                independent.append(['d', value.isoformat()])
            elif value is None:
                independent.append(['z'])
            else:
                independent.append(['s', value])
        cursor = con.execute('INSERT OR IGNORE INTO seen VALUES (?)', (json.dumps(independent),))
        if not cursor.rowcount:
            counts['duplicates'] += 1
            continue
        counts['retained'] += 1
        reference_limbs = [a+b for a, b in zip(reference_limbs, position_limbs(index, row_num))]
        invoice, code, customer = (identifier(values[i]) for i in (0, 1, 6))
        for col, ident in zip((0, 1, 6), (invoice, code, customer)):
            if ident:
                normalization[col][ident].add(json.dumps([type(values[col]).__name__, values[col]], default=str))
        q, p, date = values[3], values[5], values[4]
        valid_q = isinstance(q, (int, float)) and not isinstance(q, bool) and Decimal(str(q)) == Decimal(str(q)).to_integral_value()
        valid_p = isinstance(p, (int, float)) and not isinstance(p, bool)
        flags = {'missing_invoice': not invoice, 'missing_stock': not code,
                 'invalid_date': not isinstance(date, datetime), 'invalid_quantity': not valid_q,
                 'invalid_price': not valid_p, 'cancelled': bool(invoice and invoice.upper().startswith('C')),
                 'quantity_nonpositive': valid_q and q <= 0, 'price_nonpositive': valid_p and p <= 0,
                 'missing_customer': not customer, 'missing_description': values[2] is None,
                 'description_wrong_type': values[2] is not None and not isinstance(values[2], str)}
        reasons.update({key: int(value) for key, value in flags.items()})
        positive = not any(flags[key] for key in list(flags)[:8])
        if not positive:
            counts['invalid_or_nonpositive'] += 1
            continue
        counts['positive_valid'] += 1
        line_money = Decimal(str(q)) * Decimal(str(p))
        positive_amount += line_money
        if excluded(code):
            counts['excluded'] += 1
            excluded_counts[code] += 1
            continue
        counts['merchandise'] += 1
        counts['merchandise_missing_customer'] += int(not customer)
        amount += line_money
        item_units[code] += int(q)
        item_rows[code] += 1
        item_money[code] += line_money
        entry = baskets.setdefault(invoice, {'items': set(), 'time': date, 'lines': 0,
                                            'money': Decimal(0), 'times': set(), 'days': set(),
                                            'customers': set(), 'countries': set(), 'sheets': set()})
        entry['items'].add(code)
        entry['time'] = min(date, entry['time'])
        entry['lines'] += 1
        entry['money'] += line_money
        entry['times'].add(date)
        entry['days'].add(date.date())
        entry['sheets'].add(index)
        if customer:
            entry['customers'].add(customer)
        entry['countries'].add(values[7])
        if not (code[:5].isdigit() and len(code) >= 5):
            entry_code = special.setdefault(code, {'rows': 0, 'examples': []})
            entry_code['rows'] += 1
            if len(entry_code['examples']) < 3:
                entry_code['examples'].append({'sheet_index': index, 'source_row': row_num, 'description': values[2]})
        if len(examples) < 5:
            examples.append({'sheet_index': index, 'source_row': row_num, 'invoice': invoice,
                             'stock_code': code, 'quantity': str(q), 'price': str(p), 'line_value': str(line_money)})
        if counts['raw'] % 100000 == 0:
            con.commit()
            print(canonical({'audit_raw': counts['raw']}), flush=True)
    con.commit()
    con.close()
    months = {}
    sizes, item_frequency = Counter(), Counter()
    basket_hasher = hashlib.sha256()
    for invoice, entry in sorted(baskets.items()):
        items = sorted(entry['items'])
        basket_hasher.update((canonical([invoice, items, entry['time'].isoformat(), entry['lines']])+'\n').encode())
        sizes[len(items)] += 1
        item_frequency.update(items)
        month = entry['time'].strftime('%Y-%m')
        bucket = months.setdefault(month, {'invoices': 0, 'line_value_gbp': Decimal(0), 'lines': 0})
        bucket['invoices'] += 1
        bucket['lines'] += entry['lines']
        bucket['line_value_gbp'] += entry['money']
    # Numeric vs text spellings are reported; no source/customer values in receipt.
    collision_counts = {HEADERS[i]: sum(len(raw) > 1 for raw in by_id.values()) for i, by_id in normalization.items()}
    report = {'scope': 'INDEPENDENT_XLSX_AUDIT_NOT_SPARK', 'source_sha256': SOURCE_HASH,
              'parser': PARSER, 'policy': POLICY, 'counts': dict(counts), 'sheet_counts': dict(per_sheet),
              'raw_missing': dict(missing), 'raw_types': {k: dict(v) for k, v in type_counts.items()},
              'overlapping_retained_quality': dict(reasons),
              'scale': dict(scale), 'integer_digits': dict(integer_digits),
              'range': {k: [str(v) for v in vals] for k, vals in extrema.items()},
              'retained_position_checksum': [str(v) for v in reference_limbs],
              'checksum_method': 'sha256-four-uint64-limb-sums-of-sheet:row-v1',
              'excluded_positive_rows': dict(excluded_counts), 'merchandise_value_gbp': str(amount),
              'positive_value_gbp': str(positive_amount), 'basket_count': len(baskets),
              'basket_checksum': basket_hasher.hexdigest(), 'item_count': len(item_rows),
              'size_distribution': dict(sorted(sizes.items())), 'monthly': months,
              'normalization_multi_spelling_counts': collision_counts,
              'invoice_consistency': {name: sum(len(b[key]) > 1 for b in baskets.values()) for name, key in
                  [('multi_timestamp','times'), ('multi_day','days'), ('multi_customer','customers'),
                   ('multi_country','countries'), ('multi_sheet','sheets')]},
              'nonstandard_merchandise_codes': special, 'line_examples': examples,
              'elapsed_seconds': time.monotonic()-started}
    for bucket in months.values():
        bucket['line_value_gbp'] = str(bucket['line_value_gbp'])
    save(output / 'audit.json', report)
    save(output / 'item-reference.json', {code: {'rows': item_rows[code], 'quantity': item_units[code],
         'baskets': item_frequency[code], 'line_value_gbp': str(item_money[code])} for code in sorted(item_rows)})
    save(output / 'invoice-examples.json', {invoice: {'items': sorted(b['items']), 'time': b['time'].isoformat(),
         'lines': b['lines'], 'line_value_gbp': str(b['money'])} for invoice, b in list(sorted(baskets.items()))[:3]})
    print(canonical({'audit_complete': str(output), 'counts': dict(counts)}), flush=True)


def immutable_batch(collection, docs, ready=False):
    existing = {d['_id']: d for d in collection.find({'_id': {'$in': [d['_id'] for d in docs]}})}
    pending = []
    for doc in docs:
        old = existing.get(doc['_id'])
        if old is not None:
            check(old == doc, 'SOURCE_PAYLOAD_CONFLICT')
        else:
            check(not ready, 'READY_SOURCE_MISSING_ROW')
            # Full payload filter + unique _id catches concurrent different payloads.
            pending.append(UpdateOne(doc, {'$setOnInsert': doc}, upsert=True))
    if pending:
        collection.bulk_write(pending, ordered=True)


def import_source():
    started = time.monotonic()
    with MongoClient(URI) as client:
        db = client[DATABASE]
        metadata = db.datasets.find_one({'_id': SOURCE_HASH})
        ready = bool(metadata and metadata['status'] == 'READY')
        spec = {'_id': SOURCE_HASH, 'source_sha256': SOURCE_HASH, 'schema_version': SCHEMA,
                'parser_version': PARSER, 'source_url': 'https://archive.ics.uci.edu/dataset/502/online+retail+ii',
                'citation': 'Daqing Chen. Online Retail II. DOI 10.24432/C5CG6D. CC BY 4.0.',
                'sheets': [{'sheet_index': i, 'name': name, 'header': HEADERS, 'count': Int64(COUNTS[i])}
                           for i, name in enumerate(SHEETS)], 'expected_row_count': Int64(sum(COUNTS))}
        if metadata:
            check(all(metadata.get(k) == v for k, v in spec.items()), 'DATASET_METADATA_CONFLICT')
        else:
            db.datasets.insert_one(dict(spec, status='WRITING'))
        db.source_rows.create_index([('dataset_id', 1), ('sheet_index', 1), ('source_row', 1)], unique=True)
        seen, batch = Counter(), []
        for index, row_num, values in verified_rows():
            batch.append(document(index, row_num, values))
            seen[index] += 1
            if len(batch) == 1000:
                immutable_batch(db.source_rows, batch, ready)
                batch = []
                if sum(seen.values()) % 100000 == 0:
                    print(canonical({'import_verified_rows': sum(seen.values()), 'retry_ready': ready}), flush=True)
        if batch:
            immutable_batch(db.source_rows, batch, ready)
        for i, expected in enumerate(COUNTS):
            check(seen[i] == expected == db.source_rows.count_documents({'dataset_id': SOURCE_HASH, 'sheet_index': i}), 'IMPORT_COUNT')
        check(db.source_rows.count_documents({'dataset_id': SOURCE_HASH}) == sum(COUNTS), 'IMPORT_TOTAL')
        if not ready:
            check(db.datasets.update_one({'_id': SOURCE_HASH, 'status': 'WRITING'}, {'$set': {'status': 'READY'}}).modified_count == 1,
                  'IMPORT_STATE')
        sample = db.source_rows.find_one({'dataset_id': SOURCE_HASH})
        immutable_batch(db.source_rows, [sample], True)
        try:
            immutable_batch(db.source_rows, [dict(sample, tuple_key='deliberate-conflict')], True)
        except RuntimeError as exc:
            check(str(exc) == 'SOURCE_PAYLOAD_CONFLICT', 'CONFLICT_ERROR')
        else:
            raise RuntimeError('CONFLICT_NOT_REJECTED')
        report = {'status': 'PASS', 'source_sha256': SOURCE_HASH, 'row_count': sum(seen.values()),
                  'full_ready_retry': ready, 'same_payload_retry': True, 'different_payload_rejected': True,
                  'elapsed_seconds': time.monotonic()-started}
        path = OUT / f'import-{time.time_ns()}.json'
        save(path, report)
        print(canonical(report), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('action', choices=['audit', 'import'])
    parser.add_argument('--output', type=Path)
    args = parser.parse_args()
    if args.action == 'audit':
        audit(args.output or OUT / f'audit-{time.time_ns()}')
    else:
        from m2_runtime import verify_owned_mongo, writer_lock
        verify_owned_mongo()
        with writer_lock():
            import_source()
