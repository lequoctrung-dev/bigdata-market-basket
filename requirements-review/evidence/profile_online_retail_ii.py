"""Read-only dataset audit; not a project implementation or Spark benchmark."""
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
import hashlib
import json
import re
import sys
import time
import zipfile

import openpyxl

root = Path(__file__).resolve().parents[2]
source = root / 'data/raw/online-retail-ii/online_retail_II.xlsx'
output = Path(__file__).with_name('online-retail-ii-profile.json')
started = time.monotonic()


def identifier(value):
    if value is None:
        return ''
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def summarize_baskets(baskets):
    sizes = Counter(len(items) for items in baskets.values())
    count = sum(sizes.values())
    ordered = sorted(sizes.items())
    def quantile(fraction):
        target = max(1, int(count * fraction + 0.999999))
        cumulative = 0
        for size, frequency in ordered:
            cumulative += frequency
            if cumulative >= target:
                return size
        return None
    return {
        'invoices': count,
        'distinct_items': len(set().union(*baskets.values())) if baskets else 0,
        'single_item_invoices': sizes.get(1, 0),
        'multi_item_invoices': count - sizes.get(1, 0),
        'size_distribution': dict(ordered),
        'median_unique_items': quantile(0.5),
        'p95_unique_items': quantile(0.95),
        'max_unique_items': max(sizes, default=0),
    }


wb = openpyxl.load_workbook(source, read_only=True, data_only=False)
sheets = []
seen_rows = set()
duplicate_rows = 0
global_counts = Counter()
raw_codes = set()
raw_customers = set()
countries = Counter()
invoice_metadata = defaultdict(lambda: {'sheets': set(), 'dates': set(), 'customers': set(), 'countries': set()})
sales_baskets = defaultdict(set)
merch_baskets = defaultdict(set)
sales_times = defaultdict(set)
non_merch = defaultdict(lambda: {'rows': 0, 'descriptions': Counter(), 'positive_rows': 0})
names_by_code = defaultdict(set)
all_date_min = None
all_date_max = None
total_rows = 0
positive_gross = Decimal('0')
positive_rows_without_customer = 0
merch_pattern = re.compile(r'^\d{5}[A-Za-z]*$')

for ws in wb.worksheets:
    iterator = ws.iter_rows(values_only=True)
    headers = list(next(iterator))
    required = ['Invoice', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'Price', 'Customer ID', 'Country']
    if headers != required:
        raise ValueError(f'Unexpected headers in {ws.title}: {headers!r}')
    missing = Counter()
    types = {name: Counter() for name in headers}
    counts = Counter()
    sheet_seen_rows = set()
    sheet_invoices = set()
    sheet_codes = set()
    date_min = None
    date_max = None
    for row_number, values in enumerate(iterator, 2):
        if all(blank(value) for value in values):
            counts['fully_blank_rows'] += 1
            continue
        counts['data_rows'] += 1
        total_rows += 1
        repeated_in_sheet = values in sheet_seen_rows
        counts['duplicate_rows_within_sheet_after_first'] += int(repeated_in_sheet)
        counts['distinct_decoded_rows_already_seen_in_previous_sheet'] += int(values in seen_rows and not repeated_in_sheet)
        sheet_seen_rows.add(values)
        if values in seen_rows:
            counts['exact_duplicate_rows_after_first'] += 1
            duplicate_rows += 1
        else:
            seen_rows.add(values)
        for name, value in zip(headers, values):
            types[name][type(value).__name__] += 1
            if blank(value):
                missing[name] += 1
        inv_raw, code_raw, desc, quantity, date, price, customer_raw, country = values
        inv, code, customer = map(identifier, (inv_raw, code_raw, customer_raw))
        cancelled = inv.upper().startswith('C')
        counts['cancellation_rows'] += int(cancelled)
        valid_quantity = isinstance(quantity, (int, float)) and not isinstance(quantity, bool)
        valid_price = isinstance(price, (int, float)) and not isinstance(price, bool)
        counts['invalid_quantity_type_rows'] += int(not valid_quantity)
        counts['invalid_price_type_rows'] += int(not valid_price)
        if valid_quantity:
            counts['negative_quantity_rows'] += int(quantity < 0)
            counts['zero_quantity_rows'] += int(quantity == 0)
            counts['non_cancel_negative_quantity_rows'] += int(quantity < 0 and not cancelled)
        if valid_price:
            counts['negative_price_rows'] += int(price < 0)
            counts['zero_price_rows'] += int(price == 0)
        valid_date = isinstance(date, datetime)
        counts['invalid_date_rows'] += int(not valid_date)
        if valid_date:
            date_min = date if date_min is None else min(date_min, date)
            date_max = date if date_max is None else max(date_max, date)
            all_date_min = date if all_date_min is None else min(all_date_min, date)
            all_date_max = date if all_date_max is None else max(all_date_max, date)
        if inv:
            sheet_invoices.add(inv)
            meta = invoice_metadata[inv]
            meta['sheets'].add(ws.title)
            if valid_date:
                meta['dates'].add(date)
            if customer:
                meta['customers'].add(customer)
            if not blank(country):
                meta['countries'].add(country)
        if code:
            raw_codes.add(code)
            sheet_codes.add(code)
            if not blank(desc):
                names_by_code[code].add(str(desc).strip())
        if customer:
            raw_customers.add(customer)
        if not blank(country):
            countries[country] += 1
        positive = bool(inv and code and valid_date and valid_quantity and valid_price and quantity > 0 and price > 0 and not cancelled)
        counts['positive_sales_candidate_rows'] += int(positive)
        counts['positive_sales_candidate_missing_description'] += int(positive and blank(desc))
        if positive:
            sales_baskets[inv].add(code)
            sales_times[inv].add(date)
            positive_gross += Decimal(str(quantity)) * Decimal(str(price))
            positive_rows_without_customer += int(not customer)
        matched = bool(merch_pattern.fullmatch(code))
        if code and not matched:
            entry = non_merch[code]
            entry['rows'] += 1
            entry['positive_rows'] += int(positive)
            if not blank(desc):
                entry['descriptions'][str(desc).strip()] += 1
        if positive and matched:
            merch_baskets[inv].add(code)
            counts['merchandise_pattern_candidate_rows'] += 1
        if total_rows % 200000 == 0:
            print(json.dumps({'rows_scanned': total_rows, 'sheet': ws.title}), flush=True)
    global_counts.update(counts)
    sheets.append({
        'name': ws.title, 'declared_rows_including_header': ws.max_row,
        'headers': headers, 'counts': dict(counts),
        'missing_values': {name: missing[name] for name in headers},
        'cell_types': {name: dict(counter) for name, counter in types.items()},
        'distinct_invoices': len(sheet_invoices), 'distinct_stock_codes': len(sheet_codes),
        'date_min': date_min.isoformat() if date_min else None,
        'date_max': date_max.isoformat() if date_max else None,
    })
    print(json.dumps({'sheet_complete': ws.title, 'data_rows': counts['data_rows']}), flush=True)
wb.close()

report = {
    'profile_kind': 'full workbook scan, read-only; conditional candidate baskets; no FP-Growth or Spark runtime',
    'scanned_at_utc': datetime.now(timezone.utc).isoformat(),
    'source_path': source.relative_to(root).as_posix(),
    'source_bytes': source.stat().st_size,
    'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'source_url': 'https://archive.ics.uci.edu/dataset/502/online+retail+ii',
    'runtime': {'python': sys.version.split()[0], 'openpyxl': openpyxl.__version__},
    'sheets': sheets, 'total_data_rows': total_rows, 'counts': dict(global_counts),
    'exact_duplicate_rows_after_first': duplicate_rows, 'unique_decoded_rows': len(seen_rows),
    'distinct_raw_invoices': len(invoice_metadata), 'distinct_raw_stock_codes': len(raw_codes),
    'distinct_nonmissing_customer_ids': len(raw_customers), 'country_row_counts': dict(countries),
    'date_min': all_date_min.isoformat(), 'date_max': all_date_max.isoformat(),
    'invoice_key_checks': {
        'invoices_present_in_multiple_sheets': sum(len(m['sheets']) > 1 for m in invoice_metadata.values()),
        'invoices_with_multiple_timestamps': sum(len(m['dates']) > 1 for m in invoice_metadata.values()),
        'invoices_with_multiple_nonmissing_customer_ids': sum(len(m['customers']) > 1 for m in invoice_metadata.values()),
        'invoices_with_multiple_countries': sum(len(m['countries']) > 1 for m in invoice_metadata.values()),
        'positive_sales_candidate_invoices_with_multiple_timestamps': sum(len(times) > 1 for times in sales_times.values()),
        'positive_sales_candidate_invoices_with_multiple_calendar_dates': sum(len({t.date() for t in times}) > 1 for times in sales_times.values()),
    },
    'positive_sales_candidate_definition': 'nonempty invoice/code; valid datetime; non-C invoice; numeric quantity>0 and price>0; missing Customer ID/Description retained; exact duplicates retained',
    'positive_sales_candidate_baskets': summarize_baskets(sales_baskets),
    'positive_sales_candidate_gross_line_value_gbp': str(positive_gross),
    'positive_sales_candidate_missing_customer_rows': positive_rows_without_customer,
    'merchandise_pattern_candidate_definition': 'positive sales candidate, additionally StockCode matches five digits plus optional letters; a proposed diagnostic, not verified business classification',
    'merchandise_pattern_candidate_baskets': summarize_baskets(merch_baskets),
    'nonmatching_stock_codes': [
        {'code': code, 'rows': entry['rows'], 'positive_rows': entry['positive_rows'], 'common_descriptions': entry['descriptions'].most_common(3)}
        for code, entry in sorted(non_merch.items(), key=lambda p: (-p[1]['rows'], p[0]))
    ],
    'stock_codes_with_multiple_nonblank_descriptions': sum(len(names) > 1 for names in names_by_code.values()),
    'elapsed_seconds': round(time.monotonic() - started, 2),
    'limits': ['No source rows deleted or edited', 'No invoice-level cancellation matching', 'Missing IDs do not establish distinct people', 'Naive source timestamps have no verified timezone', 'No Spark/FP-Growth fit, no MongoDB/Spring integration, no two-machine verification', 'Duplicate decoded rows count is observation, not authorization to delete'],
}
output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'profile_saved': str(output.relative_to(root)), 'rows': total_rows, 'elapsed_seconds': report['elapsed_seconds']}), flush=True)
