"""Audit approved row deduplication and item exclusion; never rewrite source data."""
from collections import Counter, defaultdict
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
import hashlib
import json
import time

import openpyxl

root = Path(__file__).resolve().parents[2]
source = root / 'data/raw/online-retail-ii/online_retail_II.xlsx'
output = Path(__file__).with_name('online-retail-ii-approved-policy.json')
start = time.monotonic()
seen = set()
counts = Counter()
sheet_counts = {}
positive_baskets = defaultdict(set)
allowed_baskets = defaultdict(set)
excluded_codes = Counter()
review_codes = Counter()
positive_value = Decimal('0')
allowed_value = Decimal('0')
base_excluded = {'POST', 'DOT', 'C2', 'BANK CHARGES', 'AMAZONFEE', 'D', 'ADJUST', 'S', 'M'}
# Same observed meanings as approved adjustment/manual categories; explicitly recorded.
additional_adjustment_codes = {'ADJUST2', 'B'}


def ident(value):
    if value is None:
        return ''
    if isinstance(value, (int, float)) and float(value).is_integer():
        return str(int(value))
    return str(value).strip()


def baskets_summary(baskets):
    sizes = Counter(len(items) for items in baskets.values())
    return {
        'invoices': len(baskets),
        'distinct_items': len(set().union(*baskets.values())) if baskets else 0,
        'single_item_invoices': sizes.get(1, 0),
        'multi_item_invoices': sum(n for size, n in sizes.items() if size >= 2),
        'max_unique_items': max(sizes, default=0),
    }


wb = openpyxl.load_workbook(source, read_only=True, data_only=False)
expected_sheets = ['Year 2009-2010', 'Year 2010-2011']
assert wb.sheetnames == expected_sheets, wb.sheetnames
for sheet in wb.worksheets:
    rows = sheet.iter_rows(values_only=True)
    assert list(next(rows)) == ['Invoice', 'StockCode', 'Description', 'Quantity', 'InvoiceDate', 'Price', 'Customer ID', 'Country']
    per_sheet = Counter()
    for row_number, values in enumerate(rows, 2):
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        counts['raw_rows'] += 1
        per_sheet['raw_rows'] += 1
        if values in seen:
            counts['removed_exact_duplicates'] += 1
            per_sheet['removed_exact_duplicates'] += 1
            continue
        seen.add(values)
        counts['kept_first_occurrence_rows'] += 1
        per_sheet['kept_first_occurrence_rows'] += 1
        inv_raw, code_raw, desc, quantity, date, price, customer_raw, country = values
        inv, code, customer = ident(inv_raw), ident(code_raw), ident(customer_raw)
        valid_numeric = all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in (quantity, price))
        positive = bool(inv and code and isinstance(date, datetime) and valid_numeric and quantity > 0 and price > 0 and not inv.upper().startswith('C'))
        if not positive:
            counts['non_positive_or_invalid_rows_after_dedup'] += 1
            continue
        counts['positive_rows_after_dedup'] += 1
        counts['positive_rows_missing_customer'] += int(not customer)
        positive_baskets[inv].add(code)
        value = Decimal(str(quantity)) * Decimal(str(price))
        positive_value += value
        excluded = code.upper() in (base_excluded | additional_adjustment_codes) or code.lower().startswith('gift_') or code.upper() == 'GIFT'
        if excluded:
            counts['excluded_positive_item_rows'] += 1
            excluded_codes[code] += 1
            continue
        counts['allowed_positive_item_rows'] += 1
        allowed_baskets[inv].add(code)
        allowed_value += value
        if code in {'TEST001', 'TEST002'}:
            review_codes[code] += 1
        if counts['raw_rows'] % 200000 == 0:
            print(json.dumps({'raw_rows_scanned': counts['raw_rows']}), flush=True)
    sheet_counts[sheet.title] = dict(per_sheet)
    print(json.dumps({'sheet_complete': sheet.title, 'counts': dict(per_sheet)}), flush=True)
wb.close()
raw_hash = hashlib.sha256(source.read_bytes()).hexdigest()
original = json.loads(Path(__file__).with_name('online-retail-ii-profile.json').read_text(encoding='utf-8'))
assert raw_hash == original['source_sha256']
assert counts['raw_rows'] == original['total_data_rows']
assert counts['kept_first_occurrence_rows'] == original['unique_decoded_rows']
assert counts['removed_exact_duplicates'] == original['exact_duplicate_rows_after_first']
assert counts['raw_rows'] == counts['removed_exact_duplicates'] + counts['kept_first_occurrence_rows']
assert counts['kept_first_occurrence_rows'] == counts['positive_rows_after_dedup'] + counts['non_positive_or_invalid_rows_after_dedup']
assert counts['positive_rows_after_dedup'] == counts['allowed_positive_item_rows'] + counts['excluded_positive_item_rows']
report = {
    'checked_at_utc': datetime.now(timezone.utc).isoformat(),
    'source_path': source.relative_to(root).as_posix(), 'source_sha256': raw_hash,
    'policy': 'P-01 keep only first occurrence of equal decoded eight-column tuples, workbook sheet order then source row; no trimming or item normalization before duplicate comparison',
    'scope': 'read-only audit; no clean dataset export, database write, Spark fit or web implementation',
    'dedup_numeric_equality': 'Decoded numeric Excel values compare by value, e.g. 1 and 1.0 are equal; text, whitespace, case, datetime and null remain distinct as decoded',
    'counts': dict(counts), 'sheet_counts': sheet_counts,
    'positive_after_dedup_baskets': baskets_summary(positive_baskets),
    'allowed_item_candidate_baskets': baskets_summary(allowed_baskets),
    'positive_line_value_gbp_after_dedup': str(positive_value),
    'allowed_item_line_value_gbp_after_dedup': str(allowed_value),
    'excluded_item_positive_rows': dict(excluded_codes),
    'p02_case_handling': 'Case-insensitive exclusion classification only; stock identity is not globally uppercased. Lowercase m has Manual description.',
    'p02_additional_observed_adjustments': {'ADJUST2': 'Adjustment by Peter', 'B': 'Adjust bad debt'},
    'p02_voucher_rule': 'gift_ prefix or GIFT, observed voucher family; no blanket exclusion of alphabetic stock codes',
    'review_only_codes_still_in_candidates': dict(review_codes),
    'review_note': 'TEST001/TEST002 descriptions say test product; excluded from neither approved-policy candidates nor source without a separate recorded decision. Candidate counts are not final merchandise counts.',
    'elapsed_seconds': round(time.monotonic() - start, 2),
    'checks': 'source hash, raw/dedup counts against original full profile and mutually exclusive stage counts reconciled',
}
output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'report_saved': output.relative_to(root).as_posix(), 'counts': report['counts']}), flush=True)
